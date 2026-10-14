"""Relabel, add, re-sort and renumber DSE questions in the MC pp by-topic banks.

Question booklet and marking scheme are edited together so their Q numbers stay in step.
Usage: python fix_mc_bank.py <MC pp by-topic dir> <out dir> <spec.json>
"""
import copy
import json
import re
import sys
from pathlib import Path

import docx
from docx.shared import Inches
from docx.text.paragraph import Paragraph

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
HEADER = re.compile(r'^Q(\d+)\s+(CE|DSE) (\d{4})\s+Q\.?\s*(\d+)')


def text(el):
    return ''.join(t.text or '' for t in el.iter(W + 't'))


def set_text(container, value):
    """Put value in the first w:t under container and blank the rest (keeps run formatting)."""
    ts = list(container.iter(W + 't'))
    ts[0].text = value
    for t in ts[1:]:
        t.text = ''


def topic_files(root, topic):
    q = next(p for p in (root / 'Questions').glob(f'Topic {topic} -*v2.docx') if not p.name.startswith('~'))
    ms = next(p for p in (root / 'Marking scheme').glob(f'Topic {topic} -*.docx') if not p.name.startswith('~'))
    return q, ms


# ---------- question booklet ----------

def parse_booklet(d):
    """Split the body into fixed elements and sections of question blocks."""
    contents = d.tables[0]
    names = [r.cells[0].text for r in contents.rows[1:]]
    sec_names = {n for n in names if re.match(r'^\d+\. ', n)}
    group_names = set(names) - sec_names
    items, cur, block = [], None, None
    for el in list(d.element.body):
        t = text(el) if el.tag == W + 'p' else None
        is_sec = t in sec_names
        is_group = t in group_names and el is not contents._tbl
        is_head = t is not None and HEADER.match(t)
        if el.tag == W + 'sectPr' or is_sec or is_group:
            block = None
            if is_sec:
                cur = {'name': t, 'head': el, 'blocks': []}
                items.append(('section', cur))
            else:
                cur = None
                items.append(('fixed', el))
        elif is_head and cur is not None:
            block = [el]
            cur['blocks'].append(block)
        elif block is not None:
            block.append(el)
        else:
            items.append(('fixed', el))
    return items


def block_ref(block):
    m = HEADER.match(text(block[0]))
    return int(m.group(1)), m.group(2), int(m.group(3)), int(m.group(4))


def relabel_header(p, src, year, q):
    runs = p.findall(W + 'r')
    assert text(runs[0]).startswith('Q'), text(p)
    set_text(runs[1], f'     {src} {year}  Q{q}')


class Templates:
    def __init__(self, d, ref_doc):
        paras = [Paragraph(p, d) for p in d.element.body.iter(W + 'p')]
        heads = [p for p in paras if HEADER.match(p.text)]
        self.header = heads[0]._p
        self.option = next(p._p for p in paras if re.match(r'^A\.\t', p.text) and len(p._p.findall(W + 'r')) == 2)
        self.statement = next(p._p for p in paras if re.match(r'^\(1\)\t', p.text) and len(p._p.findall(W + 'r')) == 2)
        i = paras.index(heads[0])
        self.stem = next(p._p for p in paras[i + 1:] if p.text and '\t' not in p.text and len(p._p.findall(W + 'r')) == 1)
        src = d if d.tables[1:] else ref_doc
        self.tbl_template = src.tables[1]._tbl
        self.doc = d

    def para(self, kind, *parts):
        p = copy.deepcopy({'opt': self.option, 'stmt': self.statement, 'stem': self.stem}[kind])
        if kind == 'stem':
            set_text(p, parts[0])
        else:
            runs = p.findall(W + 'r')
            set_text(runs[0], parts[0])
            # some booklets keep the tab as a literal character at the start of the text run
            lead = '\t' if text(runs[1]).startswith('\t') else ''
            set_text(runs[1], lead + parts[1])
        return p

    def header_para(self, src, year, q):
        p = copy.deepcopy(self.header)
        runs = p.findall(W + 'r')
        set_text(runs[0], 'Q0')
        set_text(runs[1], f'     {src} {year}  Q{q}')
        return p

    def table(self, rows):
        tbl = copy.deepcopy(self.tbl_template)
        tr_t = tbl.findall(W + 'tr')[0]
        tc_t = tr_t.findall(W + 'tc')[0]
        for tr in tbl.findall(W + 'tr'):
            tbl.remove(tr)
        ncol = max(len(r) for r in rows)
        grid = tbl.find(W + 'tblGrid')
        for g in list(grid):
            grid.remove(g)
        width = 9000 // ncol
        for _ in range(ncol):
            g = grid.makeelement(W + 'gridCol', {W + 'w': str(width)})
            grid.append(g)
        for r in rows:
            tr = copy.deepcopy(tr_t)
            for tc in tr.findall(W + 'tc'):
                tr.remove(tc)
            for value in r:
                tc = copy.deepcopy(tc_t)
                tc.find(W + 'tcPr').find(W + 'tcW').set(W + 'w', str(width))
                ps = tc.findall(W + 'p')
                for extra in ps[1:]:
                    tc.remove(extra)
                runs = ps[0].findall(W + 'r')
                for extra in runs[1:]:
                    ps[0].remove(extra)
                if runs and list(runs[0].iter(W + 't')):
                    set_text(runs[0], value)
                else:
                    for r0 in runs[:1]:
                        ps[0].remove(r0)
                    r_el = ps[0].makeelement(W + 'r', {})
                    t_el = r_el.makeelement(W + 't', {})
                    t_el.text = value
                    t_el.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                    r_el.append(t_el)
                    ps[0].append(r_el)
                tr.append(tc)
            tbl.append(tr)
        return tbl

    def picture(self, path, width_in):
        self.doc.add_picture(str(path), width=Inches(width_in))
        p = self.doc.paragraphs[-1]._p
        p.getparent().remove(p)
        pPr = p.get_or_add_pPr()
        for tag in ('keepNext', 'keepLines'):
            pPr.append(pPr.makeelement(W + tag, {}))
        jc = pPr.makeelement(W + 'jc', {W + 'val': 'center'})
        pPr.append(jc)
        return p


def build_block(tpl, q, img_dir):
    out = [tpl.header_para('DSE', q['year'], q['q'])]
    for part in q['body']:
        kind = part[0]
        if kind == 'stem':
            out.append(tpl.para('stem', part[1]))
        elif kind == 'stmt':
            out.append(tpl.para('stmt', part[1], part[2]))
        elif kind == 'table':
            out.append(tpl.table(part[1]))
        elif kind == 'image':
            out.append(tpl.picture(img_dir / part[1], part[2]))
    for letter, opt in zip('ABCD', q['options']):
        out.append(tpl.para('opt', f'{letter}.', opt))
    return out


def renumber(items):
    n = 0
    order = []
    for kind, sec in items:
        if kind != 'section':
            continue
        sec['blocks'].sort(key=lambda b: b[1])
        for blk in sec['blocks']:
            n += 1
            set_text(blk[0][0].findall(W + 'r')[0], f'Q{n}')
            order.append((blk[2], n))
    return n, order


def update_contents(d, items, total):
    counts, start = {}, 1
    for kind, sec in items:
        if kind == 'section':
            k = len(sec['blocks'])
            counts[sec['name']] = (k, start)
            start += k
    for r in d.tables[0].rows[1:]:
        name = r.cells[0].text
        if name in counts:
            k, s = counts[name]
            set_text(r.cells[1]._tc, f'{k} question' + ('s' if k != 1 else '') if k else 'no past-paper questions')
            set_text(r.cells[2]._tc, (f'Q{s}–Q{s + k - 1}' if k > 1 else f'Q{s}') if k else '—')
    for p in d.paragraphs:
        if re.match(r'^Total: \d+ questions', p.text):
            for t in p._p.iter(W + 't'):
                if re.search(r'Total: \d+', t.text or ''):
                    t.text = re.sub(r'Total: \d+', f'Total: {total}', t.text)


# ---------- marking scheme ----------

def fix_marking_scheme(ms_doc, order, relabels, new_rows):
    """order: list of (old_id, new_no); old_id is int for existing questions, ('new', i) for added ones."""
    tbl = ms_doc.tables[0]
    rows = list(tbl._tbl.findall(W + 'tr'))
    by_old = {}
    template = None
    for tr in rows:
        cells = [text(tc) for tc in tr.findall(W + 'tc')]
        m = re.match(r'^Q(\d+)$', cells[0])
        if m and len(cells) == 3:
            by_old[int(m.group(1))] = tr
            template = template or tr
    # rows are re-emitted in the booklet's new order, after each section heading row
    seq = dict((new, old) for old, new in order)
    heading_rows = [tr for tr in rows if tr not in by_old.values()]
    for tr in by_old.values():
        tbl._tbl.remove(tr)
    new_no_by_section = []
    anchor_iter = iter(order)
    # rebuild: after each heading row, insert that section's rows; booklet order already groups by section
    section_sizes = new_rows['__sections__']
    pos = 0
    for tr in heading_rows:
        name = text(tr.findall(W + 'tc')[0])
        k = section_sizes.get(name)
        if k is None:
            continue
        anchor = tr
        for new_no in range(pos + 1, pos + k + 1):
            old = seq[new_no]
            if isinstance(old, int):
                row = by_old[old]
                ref = text(row.findall(W + 'tc')[1]).strip()
                if old in relabels:
                    src, year, q, ans = relabels[old]
                    set_text(row.findall(W + 'tc')[1], f'{src} {year}  Q{q}')
                    if ans:
                        set_text(row.findall(W + 'tc')[2], ans)
            else:
                row = copy.deepcopy(template)
                spec = new_rows[old[1]]
                set_text(row.findall(W + 'tc')[1], f"DSE {spec['year']}  Q{spec['q']}")
                set_text(row.findall(W + 'tc')[2], spec['key'])
            set_text(row.findall(W + 'tc')[0], f'Q{new_no}')
            anchor.addnext(row)
            anchor = row
        pos += k
    assert pos == len(order), (pos, len(order))


def main():
    root, out, spec_path = Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3])
    spec = json.loads(spec_path.read_text())
    img_dir = spec_path.parent
    ref_doc = docx.Document(topic_files(root, 'I')[0])
    report = []
    for topic, work in spec['topics'].items():
        qf, msf = topic_files(root, topic)
        d = docx.Document(qf)
        items = parse_booklet(d)
        sections = [s for k, s in items if k == 'section']
        tpl = Templates(d, ref_doc)
        relabels = {}
        for blk_list in (s['blocks'] for s in sections):
            for i, blk in enumerate(blk_list):
                old, src, year, q = block_ref(blk)
                r = work.get('relabel', {}).get(str(old))
                if r:
                    relabel_header(blk[0], 'DSE', r['year'], r['q'])
                    relabels[old] = ('DSE', r['year'], r['q'], r.get('key'))
                    report.append(f'{topic} Q{old}: {src} {year} Q{q} -> DSE {r["year"]} Q{r["q"]}' + (f' (key -> {r["key"]})' if r.get('key') else ''))
                    src, year, q = 'DSE', r['year'], r['q']
                blk_list[i] = (blk, (src == 'DSE', year, q, old), old)
        new_rows = {}
        for i, nq in enumerate(work.get('add', [])):
            sec = sections[nq['section'] - 1] if topic != 'B + Topic D' else sections[nq['section'] - 1]
            assert sec['name'].startswith(f"{nq['section']}. "), (sec['name'], nq)
            blk = build_block(tpl, nq, img_dir)
            sec['blocks'].append((blk, (True, nq['year'], nq['q'], 10 ** 6), ('new', i)))
            new_rows[i] = nq
            report.append(f"{topic} + DSE {nq['year']} Q{nq['q']} -> {sec['name']} (key {nq['key']})")
        total, order = renumber(items)
        body = d.element.body
        for el in list(body):
            body.remove(el)
        for kind, val in items:
            if kind == 'fixed':
                body.append(val)
            else:
                body.append(val['head'])
                for blk, _, _ in val['blocks']:
                    for el in blk:
                        body.append(el)
        update_contents(d, items, total)
        new_rows['__sections__'] = {s['name']: len(s['blocks']) for s in sections}
        ms = docx.Document(msf)
        fix_marking_scheme(ms, order, relabels, new_rows)
        (out / 'Questions').mkdir(parents=True, exist_ok=True)
        (out / 'Marking scheme').mkdir(parents=True, exist_ok=True)
        d.save(out / 'Questions' / qf.name)
        ms.save(out / 'Marking scheme' / msf.name)
        report.append(f'{topic}: {total} questions')
    print('\n'.join(report))


if __name__ == '__main__':
    main()
