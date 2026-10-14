"""Build Elective 1 / Elective 2 LQ question booklets and marking schemes in the house format.

Usage: python build_elective.py <LQ pp by-topic dir> <fig dir> <out dir>
Paragraph formats are deep-copied from the Topic C booklet so the new files match the existing bank.
"""
import copy
import re
import sys
from pathlib import Path

import docx
from docx.shared import Inches

sys.path.insert(0, str(Path(__file__).parent))
import data  # noqa: E402

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
LINES_PER_MARK = 3


def text(el):
    return ''.join(t.text or '' for t in el.iter(W + 't'))


def set_text(el, value):
    ts = list(el.iter(W + 't'))
    ts[0].text = value
    ts[0].set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    for t in ts[1:]:
        t.text = ''


class Kit:
    """Template elements lifted from an existing booklet."""

    def __init__(self, path, is_scheme):
        self.doc = docx.Document(path)
        body = list(self.doc.element.body)
        find = lambda pred: next(e for e in body if pred(e))
        self.title = body[0]
        self.subtitle = body[1]
        self.name_line = None if is_scheme else body[2]
        self.contents_head = find(lambda e: text(e) == 'Contents')
        self.contents_tbl = find(lambda e: e.tag == W + 'tbl')
        self.total_line = find(lambda e: text(e).startswith('Total:'))
        self.page_break = find(lambda e: e.xpath('.//w:br[@w:type="page"]'))
        self.topic_head = find(lambda e: e.tag == W + 'p' and text(e).startswith('Topic C:') and e is not body[0])
        self.section_head = find(lambda e: text(e) == '1. Demand')
        self.q_header = find(lambda e: re.match(r'^Q1\s', text(e)))
        i = body.index(self.q_header)
        self.para = body[i + 1]
        self.picture_para = find(lambda e: e.tag == W + 'p' and e.xpath('.//w:drawing'))
        if not is_scheme:
            self.line = body[i + 2]
            self.mark_box = find(lambda e: re.match(r'^\[\s+_+ / \d+ \]$', text(e)))
            self.total_box = find(lambda e: text(e).startswith('Total  ['))
        self.sectPr = body[-1]

    def p(self, value, template=None):
        e = copy.deepcopy(template if template is not None else self.para)
        set_text(e, value)
        return e

    def header(self, n, ref, marks):
        e = copy.deepcopy(self.q_header)
        runs = e.findall(W + 'r')
        set_text(runs[0], f'Q{n}')
        set_text(runs[1], f'     {ref}')
        set_text(runs[2], f'({marks} mark{"s" if marks != 1 else ""})')
        return e

    def picture(self, path, width):
        self.doc.add_picture(str(path), width=Inches(width))
        p = self.doc.paragraphs[-1]._p
        p.getparent().remove(p)
        p.insert(0, copy.deepcopy(self.picture_para.find(W + 'pPr')))
        return p

    def table(self, rows, box=False):
        t = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        tbl = t._tbl
        tbl.getparent().remove(tbl)
        tblPr = tbl.tblPr
        borders = tblPr.makeelement(W + 'tblBorders', {})
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            borders.append(borders.makeelement(W + side, {W + 'val': 'single', W + 'sz': '4', W + 'space': '0', W + 'color': '000000'}))
        tblPr.append(borders)
        tblPr.append(tblPr.makeelement(W + 'jc', {W + 'val': 'center'}))
        for r, row in zip(t.rows, rows):
            for c, value in zip(r.cells, row):
                c.text = value
                pPr = c.paragraphs[0]._p.get_or_add_pPr()
                pPr.append(pPr.makeelement(W + 'keepNext', {}))
                if box:
                    c.width = Inches(6.0)
        return tbl


def render(kit, blocks, figs, seen=None):
    out = []
    for b in blocks:
        if isinstance(b, str):
            b = ('p', b)
        key = repr(b)
        if seen is not None:
            if key in seen:
                continue
            seen.add(key)
        if b[0] == 'p':
            out.append(kit.p(b[1]))
        elif b[0] == 'img':
            out.append(kit.picture(figs / b[1], b[2]))
        elif b[0] == 'table':
            out.append(kit.table(b[1]))
        elif b[0] == 'box':
            out.append(kit.table([['\n'.join(b[1])]], box=True))
    return out


def entries_for(elective):
    """Split each question into entries: consecutive parts that share a sub-topic."""
    out = []
    for q in sorted((q for q in data.QUESTIONS if q['elective'] == elective), key=lambda q: q['year']):
        group = []
        for part in q['parts']:
            if group and group[-1]['topic'] != part['topic']:
                out.append((q, group))
                group = []
            group.append(part)
        out.append((q, group))
    return out


def ref_label(q, parts):
    labels = [p['label'] for p in parts]
    return f"DSE {q['year']}  Q{q['q']}" + ', '.join(labels).replace('), (', '), (')


def build(kind, elective, lq_dir, figs, out_dir):
    is_scheme = kind == 'scheme'
    src = lq_dir / ('Marking scheme/Topic C - Market and Price - Marking Scheme (CE+DSE) - formatted.docx' if is_scheme
                    else 'Questions/Topic C - Market and Price - LQ Bank (CE+DSE) - formatted.docx')
    kit = Kit(src, is_scheme)
    name, subs = data.SUBTOPICS[elective]
    entries = entries_for(elective)
    by_sub = {code: [e for e in entries if e[1][0]['topic'] == code] for code, _ in subs}

    body = kit.doc.element.body
    for el in list(body):
        body.remove(el)
    body.append(kit.sectPr)
    add = kit.sectPr.addprevious

    add(kit.p(f'{name} — ' + ('Marking Scheme' if is_scheme else 'Long Question Bank'), kit.title))
    add(kit.p('HKDSE Economics Paper 2, Section C (elective) — past-paper long questions, sorted by syllabus section. '
              'Electives were not examined in 2021–2024.' if not is_scheme else
              'HKDSE Economics Paper 2, Section C (elective) — suggested answers; question numbers match the student booklet.',
              kit.subtitle))
    if kit.name_line is not None:
        add(copy.deepcopy(kit.name_line))
    add(copy.deepcopy(kit.contents_head))

    # contents table: header row, group row, one row per sub-topic
    tbl = copy.deepcopy(kit.contents_tbl)
    rows = tbl.findall(W + 'tr')
    head, group_row, sec_row = rows[0], rows[1], rows[2]
    for r in rows:
        tbl.remove(r)
    tbl.append(head)
    g = copy.deepcopy(group_row)
    for tc in g.findall(W + 'tc'):
        set_text(tc, name)
    tbl.append(g)
    n = 1
    numbering = {}
    for code, title in subs:
        k = len(by_sub[code])
        r = copy.deepcopy(sec_row)
        tcs = r.findall(W + 'tc')
        set_text(tcs[0], title)
        set_text(tcs[1], f'{k} question' + ('s' if k != 1 else '') if k else 'no past-paper questions')
        set_text(tcs[2], (f'Q{n}–Q{n + k - 1}' if k > 1 else f'Q{n}') if k else '—')
        for e in by_sub[code]:
            numbering[id(e)] = n
            n += 1
        tbl.append(r)
    add(tbl)
    total = n - 1
    add(kit.p(f'Total: {total} questions.' + ('' if is_scheme else ' Answer on the ruled lines below each part — about 3 lines per mark, as in the real paper.'),
              kit.total_line))
    add(copy.deepcopy(kit.page_break))
    add(kit.p(name, kit.topic_head))

    for code, title in subs:
        add(kit.p(title, kit.section_head))
        for e in by_sub[code]:
            q, parts = e
            marks = sum(p['marks'] for p in parts)
            add(kit.header(numbering[id(e)], ref_label(q, parts), marks))
            seen = set()
            if not is_scheme:
                for el in render(kit, q['stem'], figs, seen):
                    add(el)
            for part in parts:
                if is_scheme:
                    for el in render(kit, part['ms'], figs):
                        add(el)
                    continue
                for el in render(kit, part['pre'], figs, seen) + render(kit, part['blocks'], figs, seen):
                    add(el)
                for _ in range(part['marks'] * LINES_PER_MARK):
                    add(copy.deepcopy(kit.line))
                add(kit.p(f'[  ______ / {part["marks"]} ]', kit.mark_box))
            if not is_scheme:
                add(kit.p(f'Total  [  ______ / {marks} ]', kit.total_box))

    short = name.split(': ', 1)[1]
    label = elective.replace('E', 'Elective ')
    sub = 'Marking scheme' if is_scheme else 'Questions'
    fname = f'{label} - {short} - ' + ('Marking Scheme (DSE) - formatted.docx' if is_scheme else 'LQ Bank (DSE) - formatted.docx')
    (out_dir / sub).mkdir(parents=True, exist_ok=True)
    for section in kit.doc.sections:
        for hf in (section.header, section.first_page_header, section.even_page_header,
                   section.footer, section.first_page_footer, section.even_page_footer):
            for t in hf._element.iter(W + 't'):
                if t.text and 'Topic C' in t.text:
                    t.text = t.text.replace('Topic C', label)
    used = set(kit.doc.element.body.xpath('.//a:blip/@r:embed'))
    part = kit.doc.part
    for rid, rel in list(part.rels.items()):
        if rel.reltype.endswith('/image') and rid not in used:
            del part.rels[rid]
    kit.doc.save(out_dir / sub / fname)
    return fname, total, {code: len(v) for code, v in by_sub.items()}


if __name__ == '__main__':
    lq, figs, out = (Path(a) for a in sys.argv[1:4])
    for el in ('E1', 'E2'):
        for kind in ('questions', 'scheme'):
            print(build(kind, el, lq, figs, out))
