import docx,re,glob,json,sys,collections
from pathlib import Path
root=Path(sys.argv[1]); K=json.load(open(sys.argv[2]))
cov=collections.defaultdict(list); problems=[]
for qf in sorted((root/'Questions').glob('Topic*v2.docx')):
    topic=qf.name.split(' - ')[0]
    msf=next((root/'Marking scheme').glob(topic+' - *.docx'))
    d=docx.Document(qf)
    book=[];sec=None;secs=collections.Counter()
    names={r.cells[0].text:(r.cells[1].text,r.cells[2].text) for r in d.tables[0].rows[1:]}
    for p in d.paragraphs:
        if p.text in names and re.match(r'^\d+\. ',p.text): sec=p.text
        m=re.match(r'^Q(\d+)\s+(CE|DSE) (\d{4})\s+Q(\d+)',p.text)
        if m: book.append((int(m.group(1)),m.group(2),int(m.group(3)),int(m.group(4)))); secs[sec]+=1
    nums=[b[0] for b in book]
    if nums!=list(range(1,len(nums)+1)): problems.append(f'{topic}: booklet numbering broken')
    ms=[]
    for r in docx.Document(msf).tables[0].rows:
        c=[x.text for x in r.cells]
        m=re.match(r'^Q(\d+)$',c[0])
        if m:
            mm=re.match(r'(CE|DSE) (\d{4})\s+Q(\d+)',c[1].strip())
            ms.append((int(m.group(1)),mm.group(1),int(mm.group(2)),int(mm.group(3)),c[2].strip()))
    if [b for b in book]!=[x[:4] for x in ms]:
        diffs=[(a,b) for a,b in zip(book,[x[:4] for x in ms]) if a!=b][:3]
        problems.append(f'{topic}: booklet vs scheme mismatch {len(book)} vs {len(ms)} {diffs}')
    # contents
    start=1
    for name,(cnt,rng) in names.items():
        if not re.match(r'^\d+\. ',name): continue
        k=secs.get(name,0)
        exp=(f'{k} question'+('s' if k!=1 else '')) if k else 'no past-paper questions'
        if cnt!=exp: problems.append(f'{topic} contents {name}: {cnt} != {exp}')
        start+=k
    tot=[p.text for p in d.paragraphs if p.text.startswith('Total:')][0]
    if not tot.startswith(f'Total: {len(book)} '): problems.append(f'{topic}: {tot[:30]} vs {len(book)}')
    for n,src,y,q,a in ms:
        if src=='DSE':
            cov[y].append(q); off=K[str(y)][q-1]
            if not (a==off or (off=='*' and 'DELET' in a.upper())): problems.append(f'{topic} Q{n} DSE {y} Q{q}: {a} vs {off}')
    print(topic, len(book), 'questions ok' )
for y in sorted(cov):
    n=len(K[str(y)]); c=collections.Counter(cov[y])
    miss=[q for q in range(1,n+1) if q not in c]; dup=[q for q,v in c.items() if v>1]
    print(y,f'{len(set(cov[y]))}/{n}','missing',miss,'multi-topic',dup)
print('PROBLEMS:' ,problems or 'none')
