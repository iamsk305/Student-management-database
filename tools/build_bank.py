"""Export the Word question banks to app/data/bank.json.

The Word booklets stay the master copy: edit them, then re-run
    python3 tools/build_bank.py "<Economic Resources dir>"
and redeploy. Only DSE questions are exported (CE stays in the booklets).
"""
import json
import re
import sys
from datetime import date
from pathlib import Path

import docx

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'app' / 'data' / 'bank.json'
KEYS = json.loads((ROOT / 'data' / 'dse-mc-keys.json').read_text())

TOPICS = [
    ('A', 'Basic Economic Concepts', '基本經濟概念'),
    ('B', 'Firms and Production', '廠商與生產'),
    ('C', 'Market and Price', '市場與價格'),
    ('D', 'Competition and Market Structure', '競爭與市場結構'),
    ('E', 'Efficiency, Equity and the Role of Government', '效率、公平與政府的角色'),
    ('F', 'Measurement of Economic Performance', '經濟表現的量度'),
    ('G', 'National Income Determination and Price Level', '國民收入決定與價格水平'),
    ('H', 'Money and Banking', '貨幣與銀行'),
    ('I', 'Macroeconomic Problems and Policies', '宏觀經濟問題與政策'),
    ('J', 'International Trade and Finance', '國際貿易與金融'),
    ('E1', 'Elective 1: Monopoly Pricing and Antitrust Law', '選修一：壟斷定價與反壟斷法'),
    ('E2', 'Elective 2: Extension of Trade Theory, Economic Growth and Development', '選修二：貿易理論延伸、經濟增長與發展'),
]

# sub-topic headings exactly as the booklets print them -> Chinese name
ZH = {
    'A': ['稀少性', '免費物品與經濟物品', '機會成本', '利息', '三個基本經濟問題', '市場經濟與計劃經濟', '私有產權', '實證與規範陳述'],
    'B': ['廠商的擁有權', '有限責任', '股票與債券', '生產類型／階段', '公共物品與私有物品', '消費品與生產品', '分工', '生產要素',
          '生產與成本（短期及長期）', '規模經濟與規模不經濟', '擴張與整合', '廠商的目標'],
    'C': ['需求', '供應', '市場均衡與價格決定', '消費者盈餘與生產者盈餘', '價格的功能', '需求與供應的價格彈性',
          '市場干預（價格干預）', '市場干預（數量干預）', '市場干預（從量稅與津貼）'],
    'D': ['市場結構'],
    'E': ['效率（條件與偏離效率）', '效率（外部性）', '公平'],
    'F': ['國民收入', '一般價格水平'],
    'G': ['總需求', '短期總供應', '長期總供應', '產出與價格水平的決定'],
    'H': ['貨幣', '銀行', '貨幣供應', '貨幣需求', '貨幣市場中利率的決定'],
    'I': ['經濟周期', '通脹與通縮', '貨幣數量論', '失業', '稅收', '公共開支', '財政政策', '貨幣政策'],
    'J': ['國際貿易', '貿易壁壘', '國際收支', '匯率'],
    'E1': ['壟斷定價與效率', '價格歧視', '反競爭行為與競爭政策'],
    'E2': ['比較優勢與貿易得益', '經濟增長與發展'],
}

# Order of teaching/Topic A.docx, plus the electives at the end
TEACHING = (['A.1', 'A.2', 'A.3', 'A.4', 'A.5', 'A.6', 'A.7', 'A.8', 'C.1', 'C.2', 'C.3', 'C.5', 'C.6', 'C.7', 'C.8', 'C.9']
            + [f'B.{i}' for i in range(1, 13)] + ['D.1', 'C.4', 'E.1', 'E.2', 'E.3', 'F.1', 'F.2', 'I.1', 'I.2', 'I.4',
                                                  'G.1', 'G.2', 'G.3', 'G.4', 'I.5', 'I.6', 'I.7', 'H.1', 'H.2', 'H.3', 'H.4', 'H.5',
                                                  'I.8', 'I.3', 'J.1', 'J.2', 'J.3', 'J.4', 'E1.1', 'E1.2', 'E1.3', 'E2.1', 'E2.2'])

# Paper 2 section totals, read off each paper's cover (Section C = one elective question, 16 marks).
P2_CORE = {y: 104 for y in range(2012, 2026)} | {2024: 120}
HAS_ELECTIVE = {y: y <= 2020 or y == 2025 for y in range(2012, 2026)}

# Essay parts the booklets file under two topics (or split into halves): merged into one part, tagged with both.
MERGE = {
    (2016, 13, '(b)'): {'labels': ['(b)(i)', '(b)(ii)'], 'marks': 10},
    (2018, 11, ''): {'labels': ['(balance of payments bullet)', '(living standard, aggregate output)'], 'marks': 10},
}

HEADER = re.compile(r'^Q(\d+)\s+(CE|DSE) (\d{4})\s+Q\.?\s*(\d+)\s*(.*?)(?:\s*\((\d+) marks?\))?$')
BOX = re.compile(r'^\[\s*_+\s*/\s*(\d+)\s*\]$')
LEAD = re.compile(r'^\(([a-h])\)\s*(?:\((i{1,3}|iv|v|vi)\))?|^\((i{1,3}|iv|v|vi)\)')


def booklet_topic(path):
    """'Topic B & D - ...' -> 'B+D', 'Elective 1 - ...' -> 'E1'."""
    name = path.name
    if name.startswith('Elective'):
        return 'E' + name.split()[1]
    m = re.match(r'Topic ([A-J])(?: ?(?:&|\+) ?(?:Topic )?([A-J]))?', name)
    return m.group(1) + ('+' + m.group(2) if m.group(2) else '')


def sections(d, booklet):
    """Map each numbered heading in the contents table to its sub-topic id."""
    groups = booklet.split('+')
    out, group = {}, groups[0]
    for row in d.tables[0].rows[1:]:
        name = row.cells[0].text.strip()
        g = re.match(r'^Topic ([A-J]):', name)
        if g:
            group = g.group(1)
            continue
        m = re.match(r'^(\d+)\. (.+)$', name)
        if m:
            out[name] = (f'{group}.{m.group(1)}', m.group(2))
    return out


def iter_entries(path):
    d = docx.Document(path)
    booklet = booklet_topic(path)
    secs = sections(d, booklet)
    sub, entry = None, None
    for p in d.paragraphs:
        t = p.text.replace('\t', ' ').strip()
        if t in secs:
            sub = secs[t][0]
            continue
        h = HEADER.match(t)
        if h:
            if entry:
                yield entry
            entry = dict(src=h.group(2), year=int(h.group(3)), q=int(h.group(4)), label=h.group(5).strip(),
                         marks=int(h.group(6) or 0), sub=sub, boxes=[], booklet=booklet)
            major = minor = None
            continue
        if entry is None:
            continue
        lead = LEAD.match(t)
        if lead:
            if lead.group(1):
                major, minor = lead.group(1), lead.group(2)
            else:
                minor = lead.group(3)
        b = BOX.match(t)
        if b:
            entry['boxes'].append((int(b.group(1)), major, minor))
    if entry:
        yield entry
    return secs


def header_major(label):
    m = re.match(r'^\(([a-h])\)', label)
    return m.group(1) if m else None


def lq_parts(entry):
    """One part per mark box; an entry without boxes is a single part."""
    if not entry['boxes']:
        return [(entry['label'] or '', entry['marks'])]
    listed = re.findall(r'\(([a-h])\)', entry['label'])
    if len(listed) > 1 and len(listed) == len(entry['boxes']) and ',' in entry['label']:
        return [(f'({l})', m) for l, (m, _, _) in zip(listed, entry['boxes'])]
    parts = []
    hm = header_major(entry['label'])
    for marks, major, minor in entry['boxes']:
        major = major or hm
        label = (f'({major})' if major else '') + (f'({minor})' if minor else '')
        parts.append((label, marks))
    if len(parts) == 1 and entry['label'] and not entry['boxes'][0][1]:
        parts = [(entry['label'], parts[0][1])]
    labels = [l for l, _ in parts]
    if len(set(labels)) < len(labels):
        romans = ['i', 'ii', 'iii', 'iv', 'v', 'vi']
        parts = [(f'{l}({romans[labels[:k].count(l)]})' if labels.count(l) > 1 else l, m)
                 for k, (l, m) in enumerate(parts)]
    return parts


def mc_rows(path):
    d = docx.Document(path)
    booklet = booklet_topic(path)
    group = booklet.split('+')[0]
    sub = None
    for row in d.tables[0].rows[1:]:
        c = [x.text.strip() for x in row.cells]
        g = re.match(r'^Topic ([A-J]):', c[0])
        if g:
            group = g.group(1)
            continue
        m = re.match(r'^(\d+)\. ', c[0])
        if m and c[0] == c[1]:
            sub = f'{group}.{m.group(1)}'
            continue
        r = re.match(r'^DSE (\d{4})\s+Q(\d+)$', c[1])
        if r:
            yield int(r.group(1)), int(r.group(2)), sub, c[2]


def build(res):
    mc_dir = res / 'PP by-topic question Bank' / 'MC pp by-topic'
    lq_dir = res / 'PP by-topic question Bank' / 'LQ pp by-topic' / 'Questions'
    names = {}
    papers = {}
    problems = []

    for f in sorted((mc_dir / 'Questions').glob('Topic*v2.docx')):
        d = docx.Document(f)
        booklet = booklet_topic(f)
        for name, (sid, title) in sections(d, booklet).items():
            names[sid] = title
    for f in sorted((mc_dir / 'Marking scheme').glob('Topic*.docx')):
        if f.name.startswith('~'):
            continue
        for year, q, sub, ans in mc_rows(f):
            key = KEYS[str(year)][q - 1]
            p1 = papers.setdefault(str(year), {'p1': {}, 'p2': []})['p1']
            if q in p1:
                problems.append(f'MC {year} Q{q} tagged twice')
            p1[q] = {'n': q, 'key': None if key == '*' else key, 'sub': sub}
            if key != '*' and ans != key:
                problems.append(f'MC {year} Q{q}: booklet {ans} vs official {key}')

    for f in sorted(lq_dir.glob('*.docx')):
        if f.name.startswith('~') or not re.match(r'(Topic|Elective)', f.name):
            continue
        d = docx.Document(f)
        for name, (sid, title) in sections(d, booklet_topic(f)).items():
            names.setdefault(sid, title)
        for e in iter_entries(f):
            if e['src'] != 'DSE':
                continue
            parts = lq_parts(e)
            if sum(m for _, m in parts) != e['marks']:
                problems.append(f"LQ {e['year']} Q{e['q']}{e['label']}: parts {parts} != {e['marks']} marks")
            p2 = papers.setdefault(str(e['year']), {'p1': {}, 'p2': []})['p2']
            for label, marks in parts:
                for (my, mq, mlabel), spec in MERGE.items():
                    if (e['year'], e['q']) == (my, mq) and label in spec['labels']:
                        label, marks = mlabel, spec['marks']
                pid = f"{e['year']}-Q{e['q']}{label}"
                same = next((x for x in p2 if x['id'] == pid), None)
                if same:
                    if same['marks'] == marks:
                        if e['sub'] not in same['subs']:
                            same['subs'].append(e['sub'])
                    else:
                        problems.append(f'LQ {pid} appears twice')
                    continue
                p2.append({'id': pid, 'q': e['q'], 'label': label, 'marks': marks, 'subs': [e['sub']],
                           'elective': e['booklet'] if e['booklet'] in ('E1', 'E2') else None})

    for y, v in papers.items():
        expected = len(KEYS[y])
        if len(v['p1']) != expected:
            problems.append(f'MC {y}: {len(v["p1"])}/{expected} questions tagged')
        v['p1'] = [v['p1'][k] for k in sorted(v['p1'])]
        v['p2'].sort(key=lambda x: (x['q'], x['label']))
        core = P2_CORE[int(y)]
        tagged = sum(p['marks'] for p in v['p2'] if not p['elective'])
        if tagged > core:
            problems.append(f'LQ {y}: tagged core {tagged} > paper {core}')
        elif tagged < core:
            v['p2'].append({'id': f'{y}-other', 'q': None, 'label': 'not in topic bank', 'marks': core - tagged,
                            'subs': [], 'elective': None})
        v['p2Total'] = core + (16 if HAS_ELECTIVE[int(y)] else 0)
        v['hasElective'] = HAS_ELECTIVE[int(y)]

    topics = []
    for code, en, zh in TOPICS:
        subs = sorted((s for s in names if s.split('.')[0] == code), key=lambda s: int(s.split('.')[1]))
        topics.append({'code': code, 'en': en, 'zh': zh,
                       'subtopics': [{'id': s, 'en': names[s], 'zh': ZH[code][int(s.split('.')[1]) - 1]} for s in subs]})
    all_subs = [s['id'] for t in topics for s in t['subtopics']]
    missing = set(all_subs) - set(TEACHING)
    if missing:
        problems.append(f'not in teaching order: {sorted(missing)}')
    return {'built': date.today().isoformat(), 'topics': topics, 'teachingOrder': TEACHING,
            'papers': dict(sorted(papers.items()))}, problems


if __name__ == '__main__':
    bank, problems = build(Path(sys.argv[1]))
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(bank, ensure_ascii=False, separators=(',', ':')))
    for y, v in bank['papers'].items():
        other = next((p['marks'] for p in v['p2'] if p['q'] is None), 0)
        print(y, f"P1 {len(v['p1'])} Qs, P2 out of {v['p2Total']}, {other} marks not in topic bank")
    print('\n'.join(problems) or 'no problems')
